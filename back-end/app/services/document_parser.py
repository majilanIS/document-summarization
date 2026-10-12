
import io
from pathlib import Path

import pymupdf
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported formats: PDF, DOCX, TXT")

    if extension == ".pdf":
        try:
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        except Exception:
            raise ValueError("The PDF file is damaged or unreadable.")

    elif extension == ".docx":
        try:
            document = Document(io.BytesIO(content))
        except Exception:
            raise ValueError("The DOCX file is damaged or unreadable.")

        paragraphs = [p.text for p in document.paragraphs]

        # Text inside tables is not part of document.paragraphs.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                paragraphs.append(" | ".join(c for c in cells if c))

        text = "\n".join(paragraphs)

    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")

    text = text.strip()

    if not text:
        if extension == ".pdf":
            raise ValueError(
                "No readable text found. The PDF may be scanned."
            )
        raise ValueError("The document is empty.")

    return text
